import os
import json
import uuid
import threading
from pathlib import Path
from flask import Flask, render_template, request, jsonify, send_file
import pdfplumber
from groq import Groq
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import werkzeug.utils

app = Flask(__name__)

# Use /tmp for Render (ephemeral filesystem)
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_FOLDER = '/tmp/uploads'
OUTPUT_FOLDER = '/tmp/outputs'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['OUTPUT_FOLDER'] = OUTPUT_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB

# In-memory job store
jobs = {}

def extract_pdf_text(pdf_path):
    """Extract text from PDF using pdfplumber."""
    pages_text = []
    with pdfplumber.open(pdf_path) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            tables = page.extract_tables()
            table_text = ""
            for table in tables:
                for row in table:
                    row_clean = [str(cell or "").strip() for cell in row]
                    table_text += " | ".join(row_clean) + "\n"
            pages_text.append({
                "page": i + 1,
                "text": text.strip(),
                "tables": table_text.strip()
            })
    return pages_text

def categorize_with_groq(pages_text, api_key, filename, job_id):
    """Use Groq AI to categorize PDF content."""
    client = Groq(api_key=api_key)

    jobs[job_id]["status"] = "extracting"
    jobs[job_id]["progress"] = 20

    full_text_parts = []
    for p in pages_text:
        chunk = f"[Page {p['page']}]\n{p['text']}"
        if p['tables']:
            chunk += f"\n[Tables]\n{p['tables']}"
        full_text_parts.append(chunk)

    full_text = "\n\n".join(full_text_parts)
    if len(full_text) > 12000:
        full_text = full_text[:12000] + "\n...[truncated for processing]"

    jobs[job_id]["status"] = "analyzing"
    jobs[job_id]["progress"] = 45

    prompt = f"""Analyze this PDF document and extract structured data. Return ONLY a valid JSON object with no markdown, no explanation.

Document: "{filename}"
Content:
{full_text}

Return this exact JSON structure:
{{
  "document_summary": "2-3 sentence summary of the document",
  "document_type": "one of: Report, Invoice, Contract, Academic Paper, Manual, Letter, Resume, Presentation, Financial, Legal, Technical, Other",
  "main_topics": ["topic1", "topic2", "topic3"],
  "key_entities": [
    {{"name": "entity name", "type": "Person/Organization/Location/Date/Amount/Product", "context": "brief context"}}
  ],
  "key_facts": [
    {{"fact": "specific fact or data point", "category": "Financial/Technical/Legal/General/Date/Contact", "page": 1}}
  ],
  "action_items": ["action or recommendation if any"],
  "sentiment": "Positive/Neutral/Negative/Mixed",
  "language": "detected language",
  "estimated_date": "date if found or Unknown"
}}

Extract at least 5 key_entities and 8 key_facts if present in the document."""

    response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.1,
        max_tokens=2000
    )

    raw = response.choices[0].message.content.strip()

    if raw.startswith("```"):
        raw = raw.split("```")[1]
        if raw.startswith("json"):
            raw = raw[4:]
    raw = raw.strip().rstrip("```").strip()

    result = json.loads(raw)
    jobs[job_id]["status"] = "building_excel"
    jobs[job_id]["progress"] = 75
    return result

def build_excel(filename, pages_text, analysis, output_path):
    """Build a nicely formatted Excel workbook."""
    wb = Workbook()

    DARK_BG = "1A1A2E"
    ACCENT = "E94560"
    MID = "16213E"
    LIGHT_TEXT = "FFFFFF"
    SUBHEADER = "0F3460"
    ROW_ALT = "F0F4FF"
    ROW_NORM = "FFFFFF"

    def hdr(ws, row, col, value, bold=True, size=11, bg=SUBHEADER, fg=LIGHT_TEXT, wrap=False):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = Font(bold=bold, size=size, color=fg, name="Calibri")
        cell.fill = PatternFill("solid", start_color=bg)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
        return cell

    def val(ws, row, col, value, bold=False, size=10, bg=ROW_NORM, fg="1A1A2E", wrap=True, align="left"):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = Font(bold=bold, size=size, color=fg, name="Calibri")
        cell.fill = PatternFill("solid", start_color=bg)
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
        return cell

    thin = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    def border_range(ws, min_row, max_row, min_col, max_col):
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                ws.cell(r, c).border = thin

    # ─── Sheet 1: Summary ───────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Summary"
    ws1.sheet_view.showGridLines = False

    ws1.row_dimensions[1].height = 40
    ws1.merge_cells("A1:F1")
    title_cell = ws1.cell(1, 1, "PDF INTELLIGENCE REPORT")
    title_cell.font = Font(bold=True, size=18, color=LIGHT_TEXT, name="Calibri")
    title_cell.fill = PatternFill("solid", start_color=DARK_BG)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws1.row_dimensions[2].height = 25
    ws1.merge_cells("A2:F2")
    sub_cell = ws1.cell(2, 1, f"Source: {filename}")
    sub_cell.font = Font(bold=False, size=10, color="AAAAAA", name="Calibri")
    sub_cell.fill = PatternFill("solid", start_color=MID)
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")

    fields = [
        ("Document Type", analysis.get("document_type", "Unknown")),
        ("Language", analysis.get("language", "Unknown")),
        ("Sentiment", analysis.get("sentiment", "Unknown")),
        ("Estimated Date", analysis.get("estimated_date", "Unknown")),
        ("Total Pages", len(pages_text)),
        ("Document Summary", analysis.get("document_summary", "")),
    ]

    row = 4
    ws1.row_dimensions[3].height = 10
    for label, value in fields:
        ws1.row_dimensions[row].height = 35 if label == "Document Summary" else 22
        hdr(ws1, row, 1, label, size=10, bg=ACCENT)
        cell = ws1.cell(row=row, column=2, value=str(value))
        cell.font = Font(size=10, name="Calibri", color="1A1A2E")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.fill = PatternFill("solid", start_color=ROW_ALT if row % 2 == 0 else ROW_NORM)
        ws1.merge_cells(f"B{row}:F{row}")
        row += 1

    row += 1
    ws1.row_dimensions[row].height = 22
    ws1.merge_cells(f"A{row}:F{row}")
    hdr(ws1, row, 1, "MAIN TOPICS", size=11, bg=DARK_BG)
    row += 1
    for i, topic in enumerate(analysis.get("main_topics", [])):
        ws1.row_dimensions[row].height = 20
        val(ws1, row, 1, f"  *  {topic}", bg=ROW_ALT if i % 2 == 0 else ROW_NORM)
        ws1.merge_cells(f"A{row}:F{row}")
        row += 1

    action_items = analysis.get("action_items", [])
    if action_items:
        row += 1
        ws1.merge_cells(f"A{row}:F{row}")
        hdr(ws1, row, 1, "ACTION ITEMS / RECOMMENDATIONS", size=11, bg=DARK_BG)
        row += 1
        for i, item in enumerate(action_items):
            ws1.row_dimensions[row].height = 22
            val(ws1, row, 1, f"  > {item}", bg=ROW_ALT if i % 2 == 0 else ROW_NORM, wrap=True)
            ws1.merge_cells(f"A{row}:F{row}")
            row += 1

    ws1.column_dimensions["A"].width = 22
    for col in ["B", "C", "D", "E", "F"]:
        ws1.column_dimensions[col].width = 20

    # ─── Sheet 2: Key Entities ───────────────────────────────────────────
    ws2 = wb.create_sheet("Entities")
    ws2.sheet_view.showGridLines = False

    ws2.row_dimensions[1].height = 35
    ws2.merge_cells("A1:D1")
    c = ws2.cell(1, 1, "KEY ENTITIES")
    c.font = Font(bold=True, size=16, color=LIGHT_TEXT, name="Calibri")
    c.fill = PatternFill("solid", start_color=DARK_BG)
    c.alignment = Alignment(horizontal="center", vertical="center")

    row = 3
    for label, col in [("Name", 1), ("Type", 2), ("Context", 3)]:
        hdr(ws2, row, col, label, bg=ACCENT, size=10)
    ws2.row_dimensions[row].height = 22
    row += 1

    entities = analysis.get("key_entities", [])
    for i, ent in enumerate(entities):
        bg = ROW_ALT if i % 2 == 0 else ROW_NORM
        ws2.row_dimensions[row].height = 22
        val(ws2, row, 1, ent.get("name", ""), bg=bg, bold=True)
        val(ws2, row, 2, ent.get("type", ""), bg=bg, align="center")
        val(ws2, row, 3, ent.get("context", ""), bg=bg, wrap=True)
        row += 1

    border_range(ws2, 3, row - 1, 1, 3)
    ws2.column_dimensions["A"].width = 25
    ws2.column_dimensions["B"].width = 18
    ws2.column_dimensions["C"].width = 55

    # ─── Sheet 3: Key Facts ──────────────────────────────────────────────
    ws3 = wb.create_sheet("Key Facts")
    ws3.sheet_view.showGridLines = False

    ws3.row_dimensions[1].height = 35
    ws3.merge_cells("A1:D1")
    c = ws3.cell(1, 1, "KEY FACTS & DATA POINTS")
    c.font = Font(bold=True, size=16, color=LIGHT_TEXT, name="Calibri")
    c.fill = PatternFill("solid", start_color=DARK_BG)
    c.alignment = Alignment(horizontal="center", vertical="center")

    row = 3
    for label, col in [("Fact", 1), ("Category", 2), ("Page", 3)]:
        hdr(ws3, row, col, label, bg=ACCENT, size=10)
    ws3.row_dimensions[row].height = 22
    row += 1

    facts = analysis.get("key_facts", [])
    for i, fact in enumerate(facts):
        bg = ROW_ALT if i % 2 == 0 else ROW_NORM
        ws3.row_dimensions[row].height = 30
        val(ws3, row, 1, fact.get("fact", ""), bg=bg, wrap=True)
        val(ws3, row, 2, fact.get("category", ""), bg=bg, align="center")
        val(ws3, row, 3, fact.get("page", ""), bg=bg, align="center")
        row += 1

    border_range(ws3, 3, row - 1, 1, 3)
    ws3.column_dimensions["A"].width = 65
    ws3.column_dimensions["B"].width = 18
    ws3.column_dimensions["C"].width = 10

    # ─── Sheet 4: Raw Text ───────────────────────────────────────────────
    ws4 = wb.create_sheet("Raw Text")
    ws4.sheet_view.showGridLines = False

    ws4.row_dimensions[1].height = 35
    ws4.merge_cells("A1:C1")
    c = ws4.cell(1, 1, "RAW TEXT BY PAGE")
    c.font = Font(bold=True, size=16, color=LIGHT_TEXT, name="Calibri")
    c.fill = PatternFill("solid", start_color=DARK_BG)
    c.alignment = Alignment(horizontal="center", vertical="center")

    row = 3
    for label, col in [("Page", 1), ("Text Content", 2), ("Tables Detected", 3)]:
        hdr(ws4, row, col, label, bg=ACCENT, size=10)
    ws4.row_dimensions[row].height = 22
    row += 1

    for i, page in enumerate(pages_text):
        bg = ROW_ALT if i % 2 == 0 else ROW_NORM
        text_preview = (page["text"] or "")[:2000]
        table_preview = (page["tables"] or "")[:500]
        lines = max(1, min(15, text_preview.count('\n') + 1))
        ws4.row_dimensions[row].height = max(20, lines * 14)
        val(ws4, row, 1, page["page"], bg=bg, align="center", bold=True)
        val(ws4, row, 2, text_preview, bg=bg, wrap=True)
        val(ws4, row, 3, table_preview, bg=bg, wrap=True)
        row += 1

    border_range(ws4, 3, row - 1, 1, 3)
    ws4.column_dimensions["A"].width = 8
    ws4.column_dimensions["B"].width = 80
    ws4.column_dimensions["C"].width = 35

    wb.save(output_path)

def process_job(job_id, pdf_path, filename, api_key):
    """Background job processor."""
    try:
        jobs[job_id]["status"] = "extracting"
        jobs[job_id]["progress"] = 10

        pages_text = extract_pdf_text(pdf_path)
        jobs[job_id]["pages"] = len(pages_text)

        analysis = categorize_with_groq(pages_text, api_key, filename, job_id)

        output_filename = f"{Path(filename).stem}_categorized_{job_id[:8]}.xlsx"
        output_path = os.path.join(app.config['OUTPUT_FOLDER'], output_filename)
        build_excel(filename, pages_text, analysis, output_path)

        jobs[job_id]["status"] = "complete"
        jobs[job_id]["progress"] = 100
        jobs[job_id]["output_file"] = output_filename
        jobs[job_id]["analysis"] = analysis

    except Exception as e:
        jobs[job_id]["status"] = "error"
        jobs[job_id]["error"] = str(e)
    finally:
        if os.path.exists(pdf_path):
            os.remove(pdf_path)

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/upload", methods=["POST"])
def upload():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    f = request.files["file"]
    api_key = request.form.get("api_key", "").strip()
    if not api_key:
        return jsonify({"error": "Groq API key is required"}), 400
    if not f.filename.lower().endswith(".pdf"):
        return jsonify({"error": "Only PDF files are supported"}), 400

    job_id = str(uuid.uuid4())
    safe_name = werkzeug.utils.secure_filename(f.filename)
    pdf_path = os.path.join(app.config['UPLOAD_FOLDER'], f"{job_id}_{safe_name}")
    f.save(pdf_path)

    jobs[job_id] = {
        "status": "queued",
        "progress": 0,
        "filename": safe_name,
        "pages": 0
    }

    thread = threading.Thread(target=process_job, args=(job_id, pdf_path, safe_name, api_key))
    thread.daemon = True
    thread.start()

    return jsonify({"job_id": job_id})

@app.route("/status/<job_id>")
def status(job_id):
    job = jobs.get(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404
    return jsonify(job)

@app.route("/download/<job_id>")
def download(job_id):
    job = jobs.get(job_id)
    if not job or job.get("status") != "complete":
        return jsonify({"error": "Not ready"}), 404
    path = os.path.join(app.config['OUTPUT_FOLDER'], job["output_file"])
    return send_file(path, as_attachment=True, download_name=job["output_file"])

if __name__ == "__main__":
    app.run(debug=False, port=5050)
